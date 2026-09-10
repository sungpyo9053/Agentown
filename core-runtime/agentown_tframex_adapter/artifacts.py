"""Local, explicit artifact creation. No shell, network, macros or existing-file edits.

Callers must select the output directory; model output never selects a path.
This module is not an assertion that an agent's content is correct or researched.
"""
from __future__ import annotations

from hashlib import sha256
from io import BytesIO
import json
import math
from pathlib import Path
import re
import unicodedata
from tempfile import mkdtemp
from typing import Any


class ArtifactContractError(ValueError):
    pass


def _text(value: Any, label: str, limit: int) -> str:
    if not isinstance(value, str) or not value.strip() or len(value) > limit:
        raise ArtifactContractError(f"{label}: nonblank text of at most {limit} characters required")
    if any(ord(ch) < 32 and ch not in "\n\t" for ch in value):
        raise ArtifactContractError(f"{label}: control characters are not allowed")
    return value.strip()


def _items(value: Any, label: str, maximum: int) -> list:
    if not isinstance(value, list) or not 1 <= len(value) <= maximum:
        raise ArtifactContractError(f"{label}: 1 to {maximum} items required")
    return value


def _presentation(spec: dict) -> tuple[bytes, dict]:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    title = _text(spec.get("title"), "title", 120)
    slides = _items(spec.get("slides"), "slides", 30)
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    deck.core_properties.title = title
    for index, data in enumerate(slides, 1):
        if not isinstance(data, dict):
            raise ArtifactContractError("slide must be an object")
        heading = _text(data.get("title"), f"slide {index} title", 52)
        bullets = [_text(item, f"slide {index} bullet", 100)
                   for item in _items(data.get("bullets"), f"slide {index} bullets", 4)]
        source = data.get("sourceNote", "")
        if source:
            source = _text(source, f"slide {index} sourceNote", 160)
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string("F5F7FA")
        heading_frame = slide.shapes.add_textbox(Inches(.7), Inches(.55), Inches(11.9), Inches(1.15)).text_frame
        heading_frame.word_wrap = True
        heading_frame.text = heading
        heading_frame.paragraphs[0].font.size = Pt(30)
        heading_frame.paragraphs[0].font.bold = True
        # Independent bounded text boxes prevent one long paragraph hiding others.
        for row, content in enumerate(bullets):
            frame = slide.shapes.add_textbox(Inches(.8), Inches(1.95 + row * 1.02), Inches(11.7), Inches(.92)).text_frame
            frame.word_wrap = True
            frame.text = content
            frame.paragraphs[0].font.size = Pt(21)
        foot = slide.shapes.add_textbox(Inches(.8), Inches(6.35), Inches(11.7), Inches(.7)).text_frame
        foot.word_wrap = True
        foot.text = f"{index}/{len(slides)}" + (" · " + source if source else "")
        foot.paragraphs[0].font.size = Pt(11)
    stream = BytesIO()
    deck.save(stream)
    content = stream.getvalue()
    loaded = Presentation(BytesIO(content))
    if len(loaded.slides) != len(slides):
        raise ArtifactContractError("slide round-trip mismatch")
    return content, {"slides": len(slides), "visualReview": "REQUIRED"}


def _row_height(values: list, width: int = 28) -> float:
    def lines(value: Any) -> int:
        return sum(max(1, math.ceil(sum(2 if unicodedata.east_asian_width(c) in "WF" else 1
                                      for c in line) / (width - 2)))
                   for line in str(value if value is not None else "").split("\n"))
    height = max(34, max(map(lines, values)) * 18 + 8)
    if height > 400:
        raise ArtifactContractError("cell text is too long to display; make cells concise or split content into additional rows")
    return height


def _workbook(spec: dict) -> tuple[bytes, dict]:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    title = _text(spec.get("title"), "title", 120)
    sheets = _items(spec.get("sheets"), "sheets", 12)
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.title = title
    names = set()
    total = 0
    for data in sheets:
        if not isinstance(data, dict):
            raise ArtifactContractError("sheet must be an object")
        name = _text(data.get("name"), "sheet name", 31)
        if re.search(r"[\\/*?:\[\]]", name) or name.casefold() in names:
            raise ArtifactContractError("sheet name must be valid and unique")
        names.add(name.casefold())
        columns = [_text(value, "column heading", 120) for value in _items(data.get("columns"), "columns", 30)]
        rows = _items(data.get("rows"), "rows", 5000)
        total += len(rows) * len(columns)
        if total > 50000:
            raise ArtifactContractError("workbook exceeds 50000 data cells")
        sheet = workbook.create_sheet(name)
        for row_index, row in enumerate([columns] + rows, 1):
            if not isinstance(row, list) or len(row) != len(columns):
                raise ArtifactContractError("row width must equal column count")
            for column, value in enumerate(row, 1):
                if value is not None and not isinstance(value, (str, int, float, bool)):
                    raise ArtifactContractError("cells must be text, finite numbers, booleans or null")
                if isinstance(value, float) and not math.isfinite(value):
                    raise ArtifactContractError("non-finite cell number")
                if isinstance(value, str) and value:
                    _text(value, "cell", 2000)
                cell = sheet.cell(row_index, column, value)
                cell.font = Font(size=11)
                # Model/user-provided strings are literal, never executable formulas.
                if isinstance(value, str):
                    cell.data_type = "s"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if row_index == 1:
                    cell.font = Font(size=11, bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="17365D")
            sheet.row_dimensions[row_index].height = _row_height(row)
        for index in range(1, len(columns) + 1):
            sheet.column_dimensions[get_column_letter(index)].width = 28
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
        sheet.print_title_rows = "1:1"
    stream = BytesIO()
    workbook.save(stream)
    content = stream.getvalue()
    check = load_workbook(BytesIO(content), read_only=True, data_only=False)
    if check.sheetnames != [sheet["name"].strip() for sheet in sheets]:
        raise ArtifactContractError("sheet round-trip mismatch")
    check.close()
    return content, {"sheets": len(sheets), "dataCells": total, "formulas": "LITERAL_ONLY", "visualReview": "REQUIRED"}


_WRITERS = {
    "pptx": (_presentation, "application/vnd.openxmlformats-officedocument.presentationml.presentation"),
    "xlsx": (_workbook, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
}


def render_artifact(spec: dict, output_directory: Path) -> dict:
    """Validate/render in memory, then create a unique file under caller-owned output."""
    if not isinstance(spec, dict) or spec.get("format") not in _WRITERS:
        raise ArtifactContractError("supported artifact formats: pptx, xlsx")
    kind = spec["format"]
    writer, media_type = _WRITERS[kind]
    content, details = writer(spec)
    root = Path(output_directory).resolve(strict=True)
    if not root.is_dir():
        raise ArtifactContractError("output directory must exist")
    target = Path(mkdtemp(prefix="agentown-result-", dir=root)) / f"result.{kind}"
    with target.open("xb") as stream:
        stream.write(content)
    return {"path": str(target), "mediaType": media_type, "bytes": len(content),
            "sha256": sha256(content).hexdigest(), "validation": details,
            "contentQuality": "NOT_ASSESSED"}


def validate_artifact_output(output: dict, options: dict) -> None:
    """Pure preflight participates in the existing bounded AI-output correction loop."""
    spec = _parse_spec(output.get("artifactJson"), options.get("format"))
    _WRITERS[spec["format"]][0](spec)  # In-memory validation only; no filesystem writes.


def _parse_spec(value: Any, approved_format: str) -> dict:
    if not isinstance(value, str) or len(value) > 1_000_000:
        raise ArtifactContractError("artifactJson must be a JSON string under 1 MB")
    spec = json.loads(value)
    if not isinstance(spec, dict) or spec.get("format") != approved_format or approved_format not in _WRITERS:
        raise ArtifactContractError("artifact format differs from the approved tool configuration")
    return spec


def local_artifact_tools(package_root: Path) -> dict:
    """Only the downloaded runner registers this closure, never the hosted runtime."""
    root = Path(package_root).resolve(strict=True)

    def render(artifactJson: str, artifactFormat: str, **_: Any) -> dict:
        spec = _parse_spec(artifactJson, artifactFormat)
        destination = root / "results"
        if destination.is_symlink():
            raise ArtifactContractError("results directory must not be a symbolic link")
        destination.mkdir(exist_ok=True)
        if destination.resolve().parent != root:
            raise ArtifactContractError("results directory is outside the package")
        result = render_artifact(spec, destination)
        return {"artifactPath": result["path"], "artifactSha256": result["sha256"],
                "artifactFormat": artifactFormat, "artifactBytes": result["bytes"],
                "artifactQuality": "REVIEW_REQUIRED"}

    return {"local.artifact.render": render}
