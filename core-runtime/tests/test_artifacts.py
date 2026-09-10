import importlib.util
import asyncio
import json
import tempfile
from pathlib import Path
import unittest

from agentown_tframex_adapter.artifacts import ArtifactContractError, render_artifact, local_artifact_tools, validate_artifact_output
from agentown_tframex_adapter import AgentownTFrameXAdapter, ExecutionNotConfigured


@unittest.skipUnless(importlib.util.find_spec("pptx") and importlib.util.find_spec("openpyxl"), "optional artifact tools not installed")
class LocalArtifactTests(unittest.TestCase):
    @unittest.skipUnless(importlib.util.find_spec("docx"), "optional report tools not installed")
    def test_docx_preserves_report_structure_and_requires_visual_review(self):
        from docx import Document
        spec = {"format": "docx", "title": "대출 연장 화면 시험 결과",
                "summary": "도서관 운영팀이 다음 시험의 확인 항목을 결정하기 위한 보고서다.",
                "sections": [{"heading": "관찰 결과", "paragraphs": [
                    "회원 12명 중 9명은 도움 없이 신청했고 3명은 대출 번호 위치를 찾는 데 도움을 받았다.",
                    "처리 시간은 측정하지 않았다. 만족도 측정 여부는 자료에 없다."]}],
                "sources": []}
        result = render_artifact(spec, self.root)
        report = Document(result["path"])
        self.assertEqual(report.paragraphs[0].style.name, "Title")
        self.assertFalse(report.styles["Title"].element.xpath("./w:pPr/w:pBdr"))
        self.assertTrue(report.styles["Normal"].paragraph_format.widow_control)
        self.assertEqual(report.styles["Normal"].paragraph_format.line_spacing, 1.15)
        self.assertEqual(report.paragraphs[0].text, spec["title"])
        self.assertIn("만족도 측정 여부는 자료에 없다.", report.paragraphs[-1].text)
        self.assertEqual(result["validation"]["sections"], 1)
        self.assertEqual(result["validation"]["pageCount"], "RENDER_REQUIRED")
        self.assertEqual(result["contentQuality"], "NOT_ASSESSED")
        self.assertEqual(report.core_properties.author, "")
        long_title = "소상공인 고객 응대 생성형 AI 도입 검토 보고서"
        long_report = Document(render_artifact(spec | {"title": long_title}, self.root)["path"])
        self.assertEqual(long_report.core_properties.title, long_title)
        self.assertEqual(long_report.paragraphs[0].text.replace("\n", " "), long_title)
        self.assertGreater(len(long_report.paragraphs[0].text.splitlines()), 1)
        self.assertTrue(long_report.paragraphs[0].text.splitlines()[-1].endswith("보고서"))
        self.assertGreater(len(long_report.paragraphs[0].text.splitlines()[-1]), 5)
        with self.assertRaises(ArtifactContractError):
            render_artifact(spec | {"sections": []}, self.root)

    @unittest.skipUnless(importlib.util.find_spec("docx"), "optional report tools not installed")
    def test_bundle_contains_actual_files_and_checksums_not_claimed_download_links(self):
        from hashlib import sha256
        from zipfile import ZipFile
        spec = {"format": "bundle", "title": "검토 결과", "artifacts": [
            {"format": "docx", "title": "검토 결과", "summary": "확인된 내용을 정리한다.",
             "sections": [{"heading": "근거", "paragraphs": ["원문에 없는 수치는 확인되지 않았다."]}]},
            {"format": "pptx", "title": "검토 결과", "slides": [
                {"title": "근거", "bullets": ["원문에 없는 수치는 미확인으로 남긴다."]}]}]}
        result = render_artifact(spec, self.root)
        with ZipFile(result["path"]) as archive:
            self.assertEqual(set(archive.namelist()), {"result-1.docx", "result-2.pptx", "manifest.json"})
            manifest = json.loads(archive.read("manifest.json"))
            self.assertEqual(manifest["contentQuality"], "NOT_ASSESSED")
            for item in manifest["artifacts"]:
                self.assertEqual(item["sha256"], sha256(archive.read(item["name"])).hexdigest())
                self.assertEqual(item["bytes"], len(archive.read(item["name"])))
        before = set(self.root.iterdir())
        with self.assertRaises(ArtifactContractError):
            render_artifact(spec | {"artifacts": [spec, spec["artifacts"][0]]}, self.root)
        with self.assertRaises(ArtifactContractError):
            render_artifact(spec | {"artifacts": [spec["artifacts"][0], spec["artifacts"][0]]}, self.root)
        self.assertEqual(set(self.root.iterdir()), before)

    def test_invalid_file_spec_retries_only_the_producer_once_without_writing_files(self):
        from tframex.models.primitives import Message
        from tframex.util.llms import BaseLLMWrapper
        class Replies(BaseLLMWrapper):
            def __init__(self, outputs):
                super().__init__(model_id="fixture")
                self.outputs, self.calls = outputs, 0
            async def chat_completion(self, messages, stream=False, **kwargs):
                value = self.outputs[min(self.calls, len(self.outputs) - 1)]
                self.calls += 1
                return Message(role="assistant", content=json.dumps({"artifactJson": value}))
        fields = [{"name": "artifactJson", "type": "string", "required": True}]
        definition = {"flowName": "spec", "steps": ["producer"], "input": "{}", "finalOutputSchema": fields,
            "agents": [{"name": "producer", "systemPrompt": "Create file content", "inputSchema": [], "outputSchema": fields,
                "outputChecks": [{"validator": "local.artifact.spec", "options": {"format": "xlsx"}}]}]}
        llm = Replies(['{"format":', json.dumps(self.workbook())])
        adapter = AgentownTFrameXAdapter(llm=llm, output_validators={"local.artifact.spec": validate_artifact_output})
        result = asyncio.run(adapter.run(definition))
        self.assertIn("artifactJson", json.loads(result["final"]))
        self.assertEqual(llm.calls, 2)
        self.assertEqual(len([e for e in adapter.trace if e["kind"] == "agent_retry"]), 1)
        self.assertEqual(list(self.root.iterdir()), [])
        bad = Replies(['{"format":'])
        with self.assertRaisesRegex(RuntimeError, "Expecting value"):
            asyncio.run(AgentownTFrameXAdapter(llm=bad, output_validators={"local.artifact.spec": validate_artifact_output}).run(definition))
        self.assertEqual(bad.calls, 2)
        missing = Replies([json.dumps(self.workbook())])
        with self.assertRaises(ExecutionNotConfigured):
            asyncio.run(AgentownTFrameXAdapter(llm=missing).run(definition))
        self.assertEqual(missing.calls, 0)

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def workbook(self):
        return {"format": "xlsx", "title": "사용자 자료", "sheets": [{"name": "비교", "columns": ["항목", "값"],
            "rows": [["실제 숫자", 1200], ["명령처럼 보이는 원문", '=WEBSERVICE("https://example.com")']]}]}

    def test_xlsx_preserves_numbers_and_treats_untrusted_formulas_as_literal_text(self):
        from openpyxl import load_workbook
        result = render_artifact(self.workbook(), self.root)
        workbook = load_workbook(result["path"])
        self.addCleanup(workbook.close)
        self.assertEqual(workbook.active["B2"].value, 1200)
        self.assertEqual(workbook.active["B3"].data_type, "s")
        self.assertEqual(result["contentQuality"], "NOT_ASSESSED")

    def test_wrapped_korean_evidence_gets_visible_row_height_without_truncation(self):
        from openpyxl import load_workbook
        spec = self.workbook()
        text = "원문에 근거한 고객 의견입니다. " * 5
        spec["sheets"][0]["rows"] = [["F1", text]]
        result = render_artifact(spec, self.root)
        workbook = load_workbook(result["path"])
        self.addCleanup(workbook.close)
        self.assertEqual(workbook.active["B2"].value, text)
        self.assertGreater(workbook.active.row_dimensions[2].height, 48)
        self.assertEqual(workbook.active.page_setup.fitToWidth, 1)

    def test_repeated_runs_create_distinct_artifacts_without_overwriting(self):
        first = render_artifact(self.workbook(), self.root)
        before = Path(first["path"]).read_bytes()
        second = render_artifact(self.workbook(), self.root)
        self.assertNotEqual(first["path"], second["path"])
        self.assertEqual(Path(first["path"]).read_bytes(), before)

    def test_model_supplied_path_cannot_escape_caller_directory(self):
        spec = self.workbook() | {"path": "../../outside.xlsx"}
        result = render_artifact(spec, self.root)
        self.assertTrue(Path(result["path"]).is_relative_to(self.root.resolve()))
        self.assertEqual(Path(result["path"]).name, "result.xlsx")

    def test_pptx_round_trip_retains_slide_titles_and_source_note(self):
        from pptx import Presentation
        spec = {"format": "pptx", "title": "검토 결과", "slides": [
            {"title": "확인된 사실", "bullets": ["사용자가 제공한 F1 자료"], "sourceNote": "출처: 사용자 입력 F1"},
            {"title": "추가 확인", "bullets": ["실제 수요는 아직 검증되지 않았습니다."]}]}
        result = render_artifact(spec, self.root)
        presentation = Presentation(result["path"])
        self.assertEqual(len(presentation.slides), 2)
        text = "\n".join(shape.text for slide in presentation.slides for shape in slide.shapes if shape.has_text_frame)
        self.assertIn("사용자 입력 F1", text)
        self.assertEqual(result["validation"]["visualReview"], "REQUIRED")

    def test_overflow_is_rejected_without_silently_truncating_or_writing(self):
        with self.assertRaises(ArtifactContractError):
            render_artifact({"format": "pptx", "title": "검토", "slides": [
                {"title": "과도한 본문", "bullets": ["가" * 101]}]}, self.root)
        self.assertEqual(list(self.root.iterdir()), [])

    def test_invalid_rows_and_duplicate_sheet_names_are_rejected(self):
        for rows in [[["missing column"]], [["bad number", float("nan")]]]:
            spec = self.workbook()
            spec["sheets"][0]["rows"] = rows
            with self.assertRaises(ArtifactContractError):
                render_artifact(spec, self.root)
        spec = self.workbook()
        spec["sheets"].append(dict(spec["sheets"][0]))
        with self.assertRaises(ArtifactContractError):
            render_artifact(spec, self.root)
        self.assertEqual(list(self.root.iterdir()), [])

    def test_tool_is_local_only_and_executes_through_the_real_runtime(self):
        output = [{"name": name, "type": kind, "required": True} for name, kind in [
            ("artifactPath", "string"), ("artifactSha256", "string"), ("artifactFormat", "string"),
            ("artifactBytes", "integer"), ("artifactQuality", "string")]]
        definition = {"flowName": "artifact", "steps": ["render"], "agents": [{
            "name": "render", "kind": "tool", "toolName": "local.artifact.render", "tools": ["local.artifact.render"],
            "inputSchema": [{"name": "artifactJson", "type": "string", "required": True}], "outputSchema": output,
            "inputDefaults": {"artifactFormat": "xlsx"}}], "finalOutputSchema": output,
            "input": json.dumps({"artifactJson": json.dumps(self.workbook())})}
        with self.assertRaises(ExecutionNotConfigured):
            asyncio.run(AgentownTFrameXAdapter().run(definition))
        result = asyncio.run(AgentownTFrameXAdapter(tools=local_artifact_tools(self.root)).run(definition))
        value = json.loads(result["final"])
        self.assertTrue(Path(value["artifactPath"]).is_file())
        self.assertEqual(value["artifactQuality"], "REVIEW_REQUIRED")

    def test_tool_refuses_symlink_output_and_unapproved_format(self):
        with tempfile.TemporaryDirectory() as outside:
            (self.root / "results").symlink_to(outside, target_is_directory=True)
            tool = local_artifact_tools(self.root)["local.artifact.render"]
            with self.assertRaises(ArtifactContractError):
                tool(json.dumps(self.workbook()), "xlsx")
            self.assertEqual(list(Path(outside).iterdir()), [])
        with self.assertRaises(ArtifactContractError):
            tool(json.dumps(self.workbook()), "pptx")


if __name__ == "__main__":
    unittest.main()
